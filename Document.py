import streamlit as st
from docx import Document


# Función para generar un documento
def generar_documento(tipo_documento, nombre, asunto):
    doc = Document()

    if tipo_documento == "Solicitud":
        doc.add_heading('Solicitud', 0)
    else:
        doc.add_heading('Oficio', 0)

    doc.add_paragraph(f"Nombre: {nombre}")
    doc.add_paragraph(f"Asunto: {asunto}")

    doc.save(f'{tipo_documento}_{nombre}.docx')
    return f'{tipo_documento}_{nombre}.docx'


# Interfaz de usuario
st.title("Generador de Documentos")

nombre = st.text_input("Nombre")
asunto = st.text_input("Asunto")
tipo_documento = st.selectbox("Tipo de Documento", ["Solicitud", "Oficio"])

if st.button("Generar Documento"):
    archivo = generar_documento(tipo_documento, nombre, asunto)
    st.success(f"Documento generado: {archivo}")
    with open(archivo, "rb") as file:
        st.download_button(
            label="Descargar Documento",
            data=file,
            file_name=archivo,
            mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )

